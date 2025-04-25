import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

# Läs CSV-data
csv_path = r"C:\Users\moab1\AppData\Local\Temp\41fa37a8-1b13-4a36-a3e2-067104a43061_OOP2_final_home_assignment (1).zip.061\OOP2_final_home_assignment\exercise_1\sales_data.csv"
data = pd.read_csv(csv_path)

# Lägg till Total Sales kolumn
data['Total Sales'] = data['Count sold'] * data['Price per item']

# Exportera till Excel
excel_path = "sales_data_analysis.xlsx"
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    # Skriv Sales Data till Excel
    data.to_excel(writer, sheet_name='Sales Data', index=False)
    sales_data_sheet = writer.sheets['Sales Data']
    
    # Formatera rubriker
    for col in sales_data_sheet.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = Font(bold=True, size=14)

    # Skapa försäljning per region
    sales_by_region = data.groupby("Region")["Total Sales"].sum()
    sales_by_region.to_excel(writer, sheet_name='Region Sales')
    region_sheet = writer.sheets['Region Sales']

    # Formatera tabellen
    for col in region_sheet.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = Font(bold=True, size=14)

    # Lägg till ett diagram
    chart = BarChart()
    chart.title = "Total Sales per Region"
    data_ref = Reference(region_sheet, min_col=2, min_row=1, max_row=len(sales_by_region)+1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(Reference(region_sheet, min_col=1, min_row=2, max_row=len(sales_by_region)+1))
    region_sheet.add_chart(chart, "E5")

print(f"Excel-fil sparad: {excel_path}")
